"""Build docs/Pulse_deck.pptx from docs/DECK.md (one slide per `##`).

    python3 docs/build_deck.py [--api http://127.0.0.1:8000 --tenant <uuid> --user <uuid>]

With --api/--tenant/--user the "Monday queue" slide carries the real top of the queue pulled from
the running system (so the slide is evidence, not a mock-up); without them it says what to look at.
Needs python-pptx (the `dev` extra).
"""

from __future__ import annotations

import argparse
import json
import re
import urllib.request
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).parent
INK, MUTED, ACCENT = RGBColor(0x1B, 0x1F, 0x2A), RGBColor(0x5B, 0x63, 0x73), RGBColor(0x1F, 0x5F, 0xBF)
TIER = {"red": RGBColor(0xC6, 0x28, 0x28), "amber": RGBColor(0xB2, 0x6A, 0x00), "green": RGBColor(0x2E, 0x7D, 0x32)}


def parse(md: str) -> list[tuple[str, str]]:
    return [(m.group(1).strip(), m.group(2).strip())
            for m in re.finditer(r"^## (.+?)\n(.*?)(?=^## |\Z)", md, re.S | re.M)]


def fetch_queue(api: str, tenant: str, user: str) -> dict:
    req = urllib.request.Request(f"{api}/queue", headers={"X-Tenant-Id": tenant, "X-User-Id": user})
    with urllib.request.urlopen(req, timeout=15) as r:
        return json.load(r)


def add_text(slide, text: str, top: float, size: int = 20, color=INK, bold=False, height: float = 1.2):
    box = slide.shapes.add_textbox(Inches(0.7), Inches(top), Inches(11.9), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size, p.font.color.rgb, p.font.bold = Pt(size), color, bold
    return box


def queue_table(slide, queue: dict) -> None:
    rows = queue["rows"][:8]
    cols = ["Account", "Tier", "Risk (90d)", "Value at risk", "Renews in", "Main driver", "Δ"]
    shape = slide.shapes.add_table(len(rows) + 1, len(cols), Inches(0.7), Inches(2.3), Inches(11.9),
                                   Inches(0.4) * (len(rows) + 1))
    table = shape.table
    for j, c in enumerate(cols):
        table.cell(0, j).text = c
    for i, r in enumerate(rows, start=1):
        delta = "—" if r["delta"] is None else f"{r['delta']:+.2f}"
        vals = [r["account_name"], r["tier"], f"{r['probability']:.0%}",
                f"${r['value_at_risk_cents'] / 100:,.0f}", f"{r['days_to_renewal']:.0f} d",
                r["dominant_driver"].replace("_", " "), delta]
        for j, v in enumerate(vals):
            cell = table.cell(i, j)
            cell.text = v
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(13)
                if j == 1:
                    p.font.color.rgb, p.font.bold = TIER.get(r["tier"], INK), True
    for row in table.rows:
        for cell in row.cells:
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(13)
    add_text(slide, f"{len(queue['rows'])} accounts in this week's queue; {queue['below_cut']} more scored "
                    f"but below the capacity cut. Real data from the running system.",
             top=2.3 + 0.4 * (len(rows) + 1) + 0.2, size=14, color=MUTED)


def build(slides: list[tuple[str, str]], queue: dict | None, out: Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)
    add_text(s, "Pulse", 2.2, size=54, color=ACCENT, bold=True)
    add_text(s, "Renewal-risk prioritisation for Customer Success", 3.4, size=26)
    add_text(s, "What I decided, what I deliberately skipped, and how I know it works.", 4.3, size=18, color=MUTED)

    for title, body in slides:
        s = prs.slides.add_slide(blank)
        add_text(s, title, 0.5, size=32, color=ACCENT, bold=True, height=0.9)
        if title.split(". ", 1)[-1].startswith("The Monday queue") and queue:
            add_text(s, "A short ordered list, with the reason and the change since last cycle.", 1.4,
                     size=18, color=MUTED, height=0.8)
            queue_table(s, queue)
        else:
            add_text(s, body, 1.6, size=22, height=4.5)
    prs.save(out)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--api"), ap.add_argument("--tenant"), ap.add_argument("--user")
    ap.add_argument("--out", default=str(HERE / "Pulse_deck.pptx"))
    a = ap.parse_args()
    q = fetch_queue(a.api, a.tenant, a.user) if a.api and a.tenant and a.user else None
    build(parse((HERE / "DECK.md").read_text()), q, Path(a.out))
    print(f"wrote {a.out} ({'with live queue' if q else 'text only'})")
