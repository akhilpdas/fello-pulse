"""``docs/build_deck.py`` embeds the real top of the queue in the deck, so it must authenticate the way a real client
does now that identity headers are gone: sign in, keep the session cookie, then read the queue.

Names the production change that makes it fail: ``fetch_queue(api, workspace, email, password)`` signing in.
"""

from __future__ import annotations

import importlib.util
import json
import threading
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path

import pytest

SPEC = importlib.util.spec_from_file_location("build_deck", Path(__file__).resolve().parents[3] / "docs" / "build_deck.py")
build_deck = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(build_deck)

ROW = {"account_id": "a1", "account_name": "Castlegate Estates", "tier": "amber", "probability": 0.1756,
       "value_at_risk_cents": 169259176, "days_to_renewal": 15.0, "dominant_driver": "usage_decline", "delta": None}


class _Handler(BaseHTTPRequestHandler):
    def log_message(self, *a):  # silence
        pass

    def _send(self, code, body, headers=()):
        raw = json.dumps(body).encode()
        self.send_response(code)
        self.send_header("Content-Type", "application/json")
        for k, v in headers:
            self.send_header(k, v)
        self.end_headers()
        self.wfile.write(raw)

    def do_POST(self):
        if self.path != "/auth/login":
            return self._send(404, {})
        body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        if body["password"] != "right password here":
            return self._send(401, {"detail": "invalid credentials"})
        self._send(200, {"user": {}}, [("Set-Cookie", "pulse_session=tok123; Path=/; HttpOnly")])

    def do_GET(self):
        if self.path == "/queue" and "pulse_session=tok123" in (self.headers.get("Cookie") or ""):
            return self._send(200, {"rows": [ROW], "below_cut": 7, "as_of": "2026-11-10T00:00:00+00:00"})
        self._send(401, {"detail": "not authenticated"})


@pytest.fixture
def api():
    server = HTTPServer(("127.0.0.1", 0), _Handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{server.server_address[1]}"
    server.shutdown()


def test_the_queue_is_fetched_by_signing_in_first(api):
    q = build_deck.fetch_queue(api, "northwind", "csm00@northwind.example.com", "right password here")
    assert q["rows"][0]["account_name"] == "Castlegate Estates" and q["below_cut"] == 7


def test_a_wrong_password_is_an_error_not_an_empty_deck(api):
    with pytest.raises(Exception, match="401|invalid"):
        build_deck.fetch_queue(api, "northwind", "csm00@northwind.example.com", "wrong")


def test_the_deck_carries_the_real_queue_rows(tmp_path, api):
    q = build_deck.fetch_queue(api, "northwind", "csm00@northwind.example.com", "right password here")
    out = tmp_path / "deck.pptx"
    build_deck.build(build_deck.parse((Path(build_deck.__file__).parent / "DECK.md").read_text()), q, out)
    from pptx import Presentation
    slides = list(Presentation(out).slides)
    assert len(slides) == 11
    cells = [c.text for sl in slides for sh in sl.shapes if sh.has_table for r in sh.table.rows for c in r.cells]
    assert "Castlegate Estates" in cells and "$1,692,592" in cells
